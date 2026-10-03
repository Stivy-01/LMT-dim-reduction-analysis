from __future__ import annotations

import argparse
import shutil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.analysis.thesis_analysis import DEFAULT_INPUT, load_recap_metadata


OUTPUT_COLUMNS = [
    "mouse_id",
    "cage_id",
    "source_sheet",
    "recap_data_range",
    "rfid",
    "rfid_from_recap",
    "sex",
    "treatment",
    "genotype",
    "strain",
    "experiment_start",
    "experiment_end",
    "stress_start",
    "phase_assignment_status",
    "skip_first_night",
    "behavior_rows_present",
    "recap_stress_raw",
    "notes",
]


def _load_behavior(path: Path) -> pd.DataFrame:
    behavior = pd.read_csv(path, usecols=["mouse_id", "interval_start"])
    behavior["mouse_id"] = pd.to_numeric(behavior["mouse_id"], errors="coerce").astype("Int64")
    behavior["interval_start"] = pd.to_datetime(behavior["interval_start"], errors="coerce")
    return behavior


def _as_excel_value(value):
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def build_analysis_metadata(
    recap_path: Path,
    behavior_path: Path,
    exclude_source_sheets: set[str] | None = None,
) -> pd.DataFrame:
    behavior = _load_behavior(behavior_path)
    metadata = load_recap_metadata(recap_path, behavior, prefer_normalized=False)
    if exclude_source_sheets:
        metadata = metadata.loc[~metadata["source_sheet"].isin(exclude_source_sheets)].copy()
    behavior_ids = set(behavior["mouse_id"].dropna().astype(int))
    metadata["behavior_rows_present"] = metadata["mouse_id"].isin(behavior_ids)
    metadata["skip_first_night"] = True
    if "rfid" not in metadata.columns:
        metadata["rfid"] = metadata.get("rfid_from_recap", "")
    else:
        metadata["rfid"] = metadata["rfid"].fillna(metadata.get("rfid_from_recap", ""))
    for column in OUTPUT_COLUMNS:
        if column not in metadata.columns:
            metadata[column] = ""
    return metadata[OUTPUT_COLUMNS].sort_values(["source_sheet", "cage_id", "mouse_id"])


def write_analysis_metadata_sheet(recap_path: Path, metadata: pd.DataFrame, sheet_name: str) -> None:
    backup = recap_path.with_suffix(recap_path.suffix + ".bak")
    if not backup.exists():
        shutil.copy2(recap_path, backup)

    workbook = load_workbook(recap_path)
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    worksheet = workbook.create_sheet(sheet_name)

    worksheet.append(list(metadata.columns))
    for row in metadata.itertuples(index=False, name=None):
        worksheet.append([_as_excel_value(value) for value in row])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    date_columns = {"experiment_start", "experiment_end", "stress_start"}
    for column_index, column_name in enumerate(metadata.columns, start=1):
        letter = get_column_letter(column_index)
        values = [str(column_name)] + [
            "" if pd.isna(value) else str(value)
            for value in metadata[column_name].head(200)
        ]
        width = min(max(max(len(value) for value in values) + 2, 10), 42)
        worksheet.column_dimensions[letter].width = width
        if column_name in date_columns:
            for cell in worksheet[letter][1:]:
                cell.number_format = "yyyy-mm-dd hh:mm"
        if column_name in {"notes", "recap_stress_raw"}:
            for cell in worksheet[letter]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    workbook.save(recap_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Add/update normalized Analysis_Metadata sheet in recap workbook.")
    parser.add_argument("recap", type=Path)
    parser.add_argument("--behavior", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--sheet-name", default="Analysis_Metadata")
    parser.add_argument(
        "--include-cd1",
        action="store_true",
        help="Keep CD1 rows in Analysis_Metadata. By default they are excluded because they are out of thesis scope.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    exclude_source_sheets = set() if args.include_cd1 else {"CD1"}
    metadata = build_analysis_metadata(args.recap, args.behavior, exclude_source_sheets)
    write_analysis_metadata_sheet(args.recap, metadata, args.sheet_name)
    print(f"Wrote {len(metadata)} rows to {args.recap}::{args.sheet_name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
